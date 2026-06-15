"""
test_drive_excel.py — WS06: the flagship Excel agent's in-place editor + manifest.

Covers `drive_excel` (the openpyxl in-place editor reached via run_local_tool) with
the locked write-approval UX (*_edited.xlsx copy by default, opt-in overwrite, NOT
encrypted), and the finalized `data_analyst` manifest. The read_file/list_dir/
write_file fs tools are covered in test_remote_router.py (WS05).

Hermetic, mirrors test_deliverables.py: tmp_path, openpyxl round-trips. drive_excel
output is a user-directory file → loaded directly (NOT via crypto.read_plaintext).
"""
import os
import sys

import pytest
from openpyxl import Workbook, load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from agent_team import config, crypto, deliverables  # noqa: E402


def _book(path, sheet="Data"):
    """A small real workbook: header + two numeric rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(["A", "B"])
    ws.append([1, 2])
    ws.append([3, 4])
    wb.save(str(path))
    return str(path)


def _run(name, tool_input, roots):
    return deliverables.run_local_tool(name, tool_input, [str(r) for r in roots])


# ---------------------------------------------------------------------------
# drive_excel — default copy, overwrite, numeric, deny
# ---------------------------------------------------------------------------
def test_drive_excel_saves_edited_copy_by_default(tmp_path):
    src = _book(tmp_path / "book.xlsx")
    res = _run("drive_excel", {"path": src, "edits": [
        {"sheet": "Data", "cell": "C1", "value": "Total"},
        {"sheet": "Data", "cell": "C2", "formula": "=A2+B2"},
        {"sheet": "Data", "append_row": ["x", 10, 20]},
    ]}, [tmp_path])
    assert res["ok"] and res["kind"] == "Excel edit"
    # Default: a *_edited.xlsx COPY, original untouched.
    edited = tmp_path / "book_edited.xlsx"
    assert edited.exists()
    assert res["filename"] == "book_edited.xlsx"
    # The original is unchanged (no C1).
    assert load_workbook(src)["Data"]["C1"].value is None

    ws = load_workbook(str(edited))["Data"]            # NOT decrypted — user-openable
    assert ws["C1"].value == "Total"
    assert ws["C2"].value == "=A2+B2"                  # stored as a real formula
    assert ws["A4"].value == "x" and ws["C4"].value == 20   # appended row


def test_drive_excel_overwrite_edits_original_in_place(tmp_path):
    src = _book(tmp_path / "book.xlsx")
    res = _run("drive_excel", {"path": src, "overwrite": True,
                               "edits": [{"sheet": "Data", "cell": "C1", "value": "X"}]}, [tmp_path])
    assert res["ok"] and res["path"] == src
    assert not (tmp_path / "book_edited.xlsx").exists()   # no copy made
    assert load_workbook(src)["Data"]["C1"].value == "X"  # original edited in place


def test_drive_excel_numbers_stay_numeric(tmp_path):
    src = _book(tmp_path / "book.xlsx")
    _run("drive_excel", {"path": src, "edits": [{"cell": "B5", "value": 1234}]}, [tmp_path])
    ws = load_workbook(str(tmp_path / "book_edited.xlsx")).active
    assert isinstance(ws["B5"].value, int) and ws["B5"].value == 1234


def test_drive_excel_output_is_not_encrypted(tmp_path):
    src = _book(tmp_path / "book.xlsx")
    _run("drive_excel", {"path": src, "edits": [{"cell": "C1", "value": "ok"}]}, [tmp_path])
    # A plain openpyxl load (no crypto) must work — i.e. the file is plaintext on disk.
    load_workbook(str(tmp_path / "book_edited.xlsx"))


def test_drive_excel_denies_outside_workspace(tmp_path):
    root = tmp_path / "approved"
    root.mkdir()
    src = _book(tmp_path / "outside.xlsx")             # exists, but OUTSIDE the root
    res = _run("drive_excel", {"path": src, "edits": [{"cell": "A1", "value": 1}]}, [root])
    assert res["ok"] is False and "outside the approved workspace" in res["message"]
    # The file was not modified.
    assert load_workbook(src)["Data"]["A1"].value == "A"


def test_drive_excel_requires_edits(tmp_path):
    src = _book(tmp_path / "book.xlsx")
    assert _run("drive_excel", {"path": src, "edits": []}, [tmp_path])["ok"] is False
    assert _run("drive_excel", {"path": src}, [tmp_path])["ok"] is False


def test_drive_excel_missing_workbook_is_clean_failure(tmp_path):
    res = _run("drive_excel", {"path": str(tmp_path / "ghost.xlsx"),
                               "edits": [{"cell": "A1", "value": 1}]}, [tmp_path])
    assert res["ok"] is False   # never raises — assert_in_workspace rejects the missing path


# ---------------------------------------------------------------------------
# Manifest — data_analyst gains the local-file workflow
# ---------------------------------------------------------------------------
def test_data_analyst_manifest_and_prompt():
    agent = config.get_agent("data_analyst")
    assert {"list_dir", "read_file", "create_excel_workbook", "drive_excel"} <= set(agent["tools"])
    prompt = agent["system_prompt"]
    for token in ("list_dir", "read_file", "drive_excel"):
        assert token in prompt
    # The VERBATIM Phase-3 contract must remain the TAIL of the prompt, untouched.
    assert prompt.rstrip().endswith('UNKNOWN_TERMS: ["term1", "term2"].')


def test_drive_excel_in_specs_and_names():
    assert "drive_excel" in deliverables.LOCAL_TOOL_NAMES
    assert [s["name"] for s in deliverables.specs_for(["drive_excel", "read_file"])] \
        == ["read_file", "drive_excel"]   # spans local specs, declaration order
