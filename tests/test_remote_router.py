"""
test_remote_router.py — WS05: RemoteRouter + the local tool-execution loop.

RemoteRouter sends every model turn to the cloud /v1 gateway with a device token and
runs deliverable/filesystem tools LOCALLY via a text protocol. All hermetic: a FAKE
http_post drives the router (no real network), and NO provider/Anthropic key is set.

Conventions mirror tests/test_deliverables.py (root path insert, tmp_path,
crypto.read_plaintext to load produced files).
"""
import json
import os
import sys
from io import BytesIO

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from agent_team import config, crypto, deliverables, providers  # noqa: E402
from agent_team.providers import RemoteRouter  # noqa: E402
from agent_team.token_manager import BudgetExceededError  # noqa: E402

CLOUD = "https://api.tangent.example"
DEVTOK = "tng-dev-" + "z" * 43
MODEL = config.HAIKU


# ---------------------------------------------------------------------------
# Fakes
# ---------------------------------------------------------------------------
def _chat(content, prompt=100, completion=50, finish="stop"):
    """An OpenAI-shaped /v1 response body."""
    return {"choices": [{"index": 0, "finish_reason": finish,
                         "message": {"role": "assistant", "content": content}}],
            "usage": {"prompt_tokens": prompt, "completion_tokens": completion}}


def _tool_call(tool, inp):
    return "```tool_call\n" + json.dumps({"tool": tool, "input": inp}) + "\n```"


class FakeTransport:
    """A scriptable http_post. Each item is a dict (returned with status 200) or a
    (status, data) tuple. Records every call so tests can assert headers/url/body."""
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def __call__(self, url, headers, json_body, timeout):
        self.calls.append({"url": url, "headers": headers, "body": json_body})
        if not self.responses:
            return 200, _chat("(no more scripted responses)")
        item = self.responses.pop(0)
        return item if isinstance(item, tuple) else (200, item)


class FakeTM:
    def __init__(self, allow=True):
        self.allow = allow
        self.recorded = []

    def ensure_can_spend(self, budget):
        if not self.allow:
            raise BudgetExceededError(budget, "budget exhausted")

    def record_cost(self, budget, cny):
        self.recorded.append((budget, cny))


# ---------------------------------------------------------------------------
# 1. complete() posts to /v1 and meters
# ---------------------------------------------------------------------------
def test_complete_posts_to_v1_and_meters():
    tp = FakeTransport([_chat("hello from the cloud", prompt=200, completion=80)])
    tm = FakeTM()
    fired = []
    r = RemoteRouter(CLOUD, DEVTOK, token_manager=tm,
                     cost_listener=lambda m, i, o: fired.append((m, i, o)),
                     http_post=tp)
    out = r.complete("sys", "hi", model=MODEL, budget="agent_team")
    assert out == "hello from the cloud"

    call = tp.calls[0]
    assert call["url"].endswith("/v1/chat/completions")
    assert call["headers"]["Authorization"] == f"Bearer {DEVTOK}"
    assert call["body"]["model"] == MODEL
    assert call["body"]["messages"][0]["role"] == "system"
    assert "max_tokens" in call["body"]
    # Metered: record_cost with the ¥ for (200, 80), and the listener fired.
    assert tm.recorded == [("agent_team", config.cost_cny(MODEL, 200, 80))]
    assert fired == [(MODEL, 200, 80)]


def test_base_url_with_v1_suffix_is_normalized():
    tp = FakeTransport([_chat("ok")])
    r = RemoteRouter(CLOUD + "/v1/", DEVTOK, http_post=tp)
    r.complete("s", "u", model=MODEL)
    assert tp.calls[0]["url"] == CLOUD + "/v1/chat/completions"


# ---------------------------------------------------------------------------
# 2. No provider/Anthropic keys needed (domestic-mode parity)
# ---------------------------------------------------------------------------
def test_no_keys_required(monkeypatch):
    for var in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "DEEPSEEK_API_KEY", "QWEN_API_KEY"):
        monkeypatch.delenv(var, raising=False)
    tp = FakeTransport([_chat("works"), _chat("final, no tools")])
    r = RemoteRouter(CLOUD, DEVTOK, http_post=tp)
    assert r.complete("s", "u", model=MODEL) == "works"
    assert "final" in r.complete_with_tools("s", "u", [], lambda n, i: "x", model=MODEL)


# ---------------------------------------------------------------------------
# 3. The local tool loop builds a REAL .xlsx end-to-end
# ---------------------------------------------------------------------------
def test_local_tool_loop_builds_real_xlsx(tmp_path):
    out_dir = str(tmp_path / "out")
    sheet = {"name": "Sales", "headers": ["Region", "Total"],
             "rows": [["华东", 120], ["华南", 98]]}
    tp = FakeTransport([
        _chat(_tool_call("create_excel_workbook", {"title": "Q4", "sheets": [sheet]})),
        _chat("Done — I built the workbook with a Sales sheet."),
    ])
    tm = FakeTM()
    r = RemoteRouter(CLOUD, DEVTOK, token_manager=tm, http_post=tp)

    def executor(name, tool_input):           # the base_agent-style executor
        return deliverables.run_tool(name, tool_input, out_dir)["message"]

    specs = deliverables.specs_for(["create_excel_workbook"])
    narration = r.complete_with_tools("sys", "make a workbook", specs, executor,
                                      model=MODEL, budget="agent_team")
    assert "Done" in narration
    assert len(tm.recorded) == 2              # one record_cost per remote turn

    # A real .xlsx was written; load it back and verify a real number survived.
    files = [f for f in os.listdir(out_dir) if f.endswith(".xlsx")]
    assert len(files) == 1
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(crypto.read_plaintext(os.path.join(out_dir, files[0]))))
    ws = wb["Sales"]
    assert ws["A1"].value == "Region"          # header row
    assert ws["B2"].value == 120 and ws["B3"].value == 98   # real numbers, not strings


# ---------------------------------------------------------------------------
# 4. read_file / write_file go through the workspace chokepoint
# ---------------------------------------------------------------------------
def test_read_file_inside_workspace_returns_text(tmp_path):
    f = tmp_path / "data.csv"
    f.write_text("region,sales\n华东,120\n", encoding="utf-8")
    res = deliverables.run_local_tool("read_file", {"path": str(f)}, [str(tmp_path)])
    assert res["ok"] and "华东" in res["message"]


def test_read_file_outside_workspace_is_denied(tmp_path):
    root = tmp_path / "approved"
    root.mkdir()
    res = deliverables.run_local_tool("read_file", {"path": "/etc/hosts"}, [str(root)])
    assert res["ok"] is False
    assert "outside the approved workspace" in res["message"]

    outside = tmp_path / "secret.txt"
    outside.write_text("nope")
    esc = deliverables.run_local_tool(
        "read_file", {"path": str(root / ".." / "secret.txt")}, [str(root)])
    assert esc["ok"] is False and "outside the approved workspace" in esc["message"]


def test_write_file_writes_and_encrypts_and_denies_escape(tmp_path):
    root = tmp_path / "ws"
    root.mkdir()
    ok = deliverables.run_local_tool(
        "write_file", {"path": str(root / "note.md"), "content": "# Hi\nbody"}, [str(root)])
    assert ok["ok"] and ok["filename"] == "note.md"
    assert crypto.read_plaintext(str(root / "note.md")).decode("utf-8") == "# Hi\nbody"

    denied = deliverables.run_local_tool(
        "write_file", {"path": str(tmp_path / "escape.md"), "content": "x"}, [str(root)])
    assert denied["ok"] is False and "outside the approved workspace" in denied["message"]


def test_local_tool_specs_shape():
    names = [s["name"] for s in deliverables.LOCAL_TOOL_SPECS]
    assert names == ["read_file", "list_dir", "write_file"]
    by = {s["name"]: s for s in deliverables.LOCAL_TOOL_SPECS}
    assert by["read_file"]["input_schema"]["required"] == ["path"]
    assert by["write_file"]["input_schema"]["required"] == ["path", "content"]


# ---------------------------------------------------------------------------
# 5. max_iterations guard
# ---------------------------------------------------------------------------
def test_max_iterations_guard(tmp_path):
    # Always returns a tool_call → the loop must stop after max_iterations.
    tp = FakeTransport([_chat(_tool_call("list_dir", {"path": str(tmp_path)}))] * 20)
    r = RemoteRouter(CLOUD, DEVTOK, http_post=tp)
    out = r.complete_with_tools(
        "s", "u", deliverables.specs_for(["list_dir"]),
        lambda n, i: deliverables.run_local_tool(n, i, [str(tmp_path)])["message"],
        model=MODEL, max_iterations=3)
    assert "Stopped after reaching the tool-call limit." in out
    assert len(tp.calls) == 3


# ---------------------------------------------------------------------------
# 6. Budget cutoff — no POST is made
# ---------------------------------------------------------------------------
def test_budget_cutoff_before_post():
    tp = FakeTransport([_chat("should never be used")])
    r = RemoteRouter(CLOUD, DEVTOK, token_manager=FakeTM(allow=False), http_post=tp)
    with pytest.raises(BudgetExceededError):
        r.complete("s", "u", model=MODEL, budget="agent_team")
    assert tp.calls == []   # cut off before any network call


# ---------------------------------------------------------------------------
# 7. HTTP error mapping
# ---------------------------------------------------------------------------
def test_http_401_maps_to_repair_error():
    tp = FakeTransport([(401, {"error": {"message": "bad token"}})])
    r = RemoteRouter(CLOUD, DEVTOK, http_post=tp)
    with pytest.raises(providers.LLMError) as ei:
        r.complete("s", "u", model=MODEL)
    assert "re-pair" in str(ei.value).lower()


def test_http_402_maps_to_balance_error():
    tp = FakeTransport([(402, {"error": {"message": "no balance"}})])
    r = RemoteRouter(CLOUD, DEVTOK, http_post=tp)
    with pytest.raises(providers.LLMError) as ei:
        r.complete("s", "u", model=MODEL)
    assert "balance" in str(ei.value).lower()


# ---------------------------------------------------------------------------
# 8. Malformed tool_call resilience
# ---------------------------------------------------------------------------
def test_malformed_tool_call_does_not_crash():
    tp = FakeTransport([
        _chat("```tool_call\n{this is not valid json\n```"),
        _chat("Recovered and produced the final answer."),
    ])
    calls_seen = []
    r = RemoteRouter(CLOUD, DEVTOK, http_post=tp)
    out = r.complete_with_tools("s", "u", deliverables.specs_for(["list_dir"]),
                                lambda n, i: calls_seen.append(n) or "x", model=MODEL)
    assert "Recovered" in out
    assert calls_seen == []          # the broken call was never dispatched to a tool
    assert len(tp.calls) == 2        # error fed back, loop continued to a final answer
