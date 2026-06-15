"""
test_deliverables.py — the deliverable BUILDERS in agent_team/deliverables.py make
real, openable files (no LLM, no network).

Run from the project root:

    .venv/bin/python -m pytest tests/test_deliverables.py -v

Focus: the create_excel_workbook tool added so a "data analytics agent that
returns a real Excel file" is honest. We drive deliverables.run_tool() exactly as
base_agent does, then read the bytes back with openpyxl and assert headers, rows,
sheet order, and — crucially — that numeric cells stay REAL numbers (so SUM works).
Files are loaded through crypto.read_plaintext so the test passes whether or not
encryption-at-rest happens to be configured in the environment.
"""

import os
import sys
from io import BytesIO

import pytest
from openpyxl import load_workbook

# Make the project root importable no matter where pytest is launched from.
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from agent_team import config, crypto, deliverables  # noqa: E402


def _load(path):
    """Open a produced .xlsx, transparently decrypting if at-rest crypto is on."""
    return load_workbook(BytesIO(crypto.read_plaintext(path)))


# ---------------------------------------------------------------------------
# Registration — the tool is wired into the same machinery as the others.
# ---------------------------------------------------------------------------
def test_excel_tool_is_registered():
    assert "create_excel_workbook" in deliverables.TOOL_NAMES
    ext, kind, builder = deliverables._BUILDERS["create_excel_workbook"]
    assert (ext, kind) == ("xlsx", "Excel workbook")
    assert callable(builder)

    specs = deliverables.specs_for(["create_excel_workbook"])
    assert [s["name"] for s in specs] == ["create_excel_workbook"]
    spec = specs[0]
    assert "xlsx" in spec["description"].lower()
    assert spec["input_schema"]["required"] == ["title", "sheets"]


# ---------------------------------------------------------------------------
# The builder writes a real, multi-sheet workbook we can read back.
# ---------------------------------------------------------------------------
def test_create_excel_workbook_writes_real_xlsx(tmp_path):
    data = {
        "filename": "sales_q4",
        "title": "Q4 Sales",
        "sheets": [
            {"name": "By Region",
             "headers": ["Region", "Revenue", "Growth"],
             "rows": [["North", 1200000, 0.12],
                      ["South", 980000.5, -0.03]]},
            {"name": "Summary",
             "headers": ["Metric", "Value"],
             "rows": [["Total", 2180000.5]]},
        ],
    }
    res = deliverables.run_tool("create_excel_workbook", data, str(tmp_path))

    assert res["ok"], res
    assert res["ext"] == "xlsx"
    assert res["kind"] == "Excel workbook"
    assert res["filename"] == "sales_q4.xlsx"
    assert os.path.exists(res["path"])

    wb = _load(res["path"])
    assert wb.sheetnames == ["By Region", "Summary"]          # order preserved

    ws = wb["By Region"]
    assert [c.value for c in ws[1]] == ["Region", "Revenue", "Growth"]
    assert ws["A1"].font.bold is True                          # header row styled
    assert ws["A2"].value == "North"
    # Numbers must arrive as REAL numbers, not strings — this is the whole point.
    assert ws["B2"].value == 1200000 and isinstance(ws["B2"].value, int)
    assert ws["C3"].value == -0.03 and isinstance(ws["C3"].value, float)
    assert wb["Summary"]["B2"].value == 2180000.5


def test_empty_sheets_still_produce_a_valid_workbook(tmp_path):
    res = deliverables.run_tool("create_excel_workbook",
                                {"title": "Empty", "sheets": []}, str(tmp_path))
    assert res["ok"], res
    wb = _load(res["path"])
    assert len(wb.sheetnames) == 1                             # never a 0-sheet (invalid) file


def test_sheet_names_are_sanitized(tmp_path):
    data = {
        "title": "t",
        "sheets": [
            {"name": "a/b:c*?[]\\", "headers": ["x"], "rows": [[1]]},  # forbidden chars
            {"name": "X" * 50},                                        # too long (>31)
            {"name": "dupe"},
            {"name": "dupe"},                                          # duplicate
        ],
    }
    res = deliverables.run_tool("create_excel_workbook", data, str(tmp_path))
    assert res["ok"], res

    names = _load(res["path"]).sheetnames
    assert not any(ch in names[0] for ch in "[]:*?/\\")        # forbidden chars stripped
    assert all(len(n) <= 31 for n in names)                    # Excel's hard limit
    assert len(set(names)) == len(names)                       # uniqueness enforced
    assert names[2] == "dupe" and names[3].startswith("dupe")


def test_filename_collisions_get_suffixed(tmp_path):
    payload = {"filename": "dup", "title": "t", "sheets": [{"name": "s"}]}
    first = deliverables.run_tool("create_excel_workbook", payload, str(tmp_path))
    second = deliverables.run_tool("create_excel_workbook", payload, str(tmp_path))
    assert first["filename"] == "dup.xlsx"
    assert second["filename"] == "dup-2.xlsx"


# ---------------------------------------------------------------------------
# Config wiring — the agents that should be able to emit an .xlsx have the tool.
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("agent_id", ["data_analyst", "media_producer"])
def test_excel_tool_granted_to_agent(agent_id):
    assert "create_excel_workbook" in config.get_agent(agent_id)["tools"]
